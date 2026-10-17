"""Excel manifest.

Sheet 1 "Reconcile": the team's tracking format (A Reconcile = batch number, B Model, C Qty),
one row per batch + model, ready to paste into their sheet.
Sheet 2 "Pieces": one row per piece, in placement order.
"""

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_THIN = Side(style="thin", color="000000")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")


def _reconcile_sheet(ws, placements, items):
    ws.title = "Reconcile"
    ws.append(["Reconcile", "Model", "Qty."])
    # Model as the team tracks it: line + die label, e.g. "LUX-15-PL", "MIRROR-17-PM"
    counts = Counter((items[p.index].batch_number,
                      items[p.index].reconcile_model or f"{items[p.index].line}-{items[p.index].label}")
                     for p in placements)
    for (batch, model), qty in sorted(counts.items(), key=lambda kv: (kv[0][0], kv[0][1])):
        ws.append([int(batch) if batch.isdigit() else batch, model, qty])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Calibri", size=11, bold=cell.row == 1)
            cell.alignment = _CENTER
            cell.border = _BOX
    for col, w in zip("ABC", (18, 30, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

COLUMNS = [
    ("Sheet", lambda it, p: p.sheet),
    ("Position", lambda it, p: p.position),
    ("Label", lambda it, p: it.label),
    ("Batch Number", lambda it, p: it.batch_number),
    ("Order Number", lambda it, p: it.order_number),
    ("HP Order ID", lambda it, p: it.hp_order_id),
    ("Component Barcode", lambda it, p: it.component_barcode),
    ("Item Barcode", lambda it, p: it.item_barcode),
    ("Print Code", lambda it, p: it.print_code),
    ("Print SKU", lambda it, p: it.print_sku),
    ("SKU", lambda it, p: it.sku),
    ("Model", lambda it, p: it.model),
    ("Case Type", lambda it, p: it.case_type),
    ("Job Status", lambda it, p: it.job_status),
    ("Order Status", lambda it, p: it.order_status),
    ("Item Status", lambda it, p: it.item_status),
    ("Reprint", lambda it, p: "Yes" if it.reprint else ""),
    ("Reprint Reason", lambda it, p: it.reprint_reason),
    ("Status Updated (HP)", lambda it, p: it.status_updated),
    ("Width mm", lambda it, p: round(p.w / 72 * 25.4, 2)),
    ("Height mm", lambda it, p: round(p.h / 72 * 25.4, 2)),
    ("File URL", lambda it, p: it.fetch_url),
]


def _dark_header(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")


def write_custom_manifest(rows, path, title):
    """Custom artwork batch: one row per piece (no Pulse order data)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pieces"
    ws.append(["Sheet", "Position", "Name on sheet", "Width mm", "Height mm", "Fit to slot", "Source file"])
    _dark_header(ws)
    for r in rows:
        ws.append(r)
    for col, w in zip("ABCDEFG", (7, 9, 40, 10, 10, 10, 40)):
        ws.column_dimensions[col].width = w
    wb.properties.title = title
    wb.save(path)


def write_manifest(placements, items, path, title):
    wb = Workbook()
    _reconcile_sheet(wb.active, placements, items)
    ws = wb.create_sheet("Pieces")
    ws.append([c[0] for c in COLUMNS])
    _dark_header(ws)
    for p in sorted(placements, key=lambda p: (p.sheet, p.position)):   # Pieces always read S1, S2, …
        ws.append([fn(items[p.index], p) for _, fn in COLUMNS])
    widths = [7, 9, 10, 14, 14, 26, 18, 14, 14, 22, 26, 20, 24, 12, 13, 12, 8, 22, 18, 10, 10, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.properties.title = title
    wb.save(path)
