"""Read barcodes from pasted text, CSV or Excel uploads; build the import templates."""

import csv
import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADER = "Component Barcode"
_CODE = re.compile(r"^B[0-9A-Z]{7,15}$")
_HEADER_NAMES = {"component barcode", "component_barcode", "componentbarcode", "barcode", "item barcode"}


def parse_text(text):
    """Barcodes separated by newlines, commas, semicolons, tabs or spaces."""
    return [t for t in re.split(r"[\s,;]+", text.upper()) if t]


def _from_rows(rows):
    """Use the 'Component Barcode' column if present, else every cell that looks like a barcode."""
    rows = [[("" if v is None else str(v)).strip() for v in r] for r in rows]
    for ri, r in enumerate(rows[:10]):
        for ci, v in enumerate(r):
            if v.lower() in _HEADER_NAMES:
                return [row[ci].upper() for row in rows[ri + 1:] if ci < len(row) and row[ci]]
    return [v.upper() for r in rows for v in r if _CODE.match(v.upper())]


def parse_upload(filename, data):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return _from_rows(ws.iter_rows(values_only=True))
    if name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        return _from_rows(csv.reader(io.StringIO(text)))
    raise ValueError("Upload a .csv or .xlsx file")


def template_csv():
    return f"{HEADER}\nB433BB89C1\nB433C55789\n".encode()


def template_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Barcodes"
    ws.append([HEADER])
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F2937")
    ws.append(["B433BB89C1"])
    ws.append(["B433C55789"])
    ws.column_dimensions["A"].width = 24
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
